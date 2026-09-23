"""测试用的 docx / pdf 字节生成器。

**为什么在测试里现场生成，而不是往仓库里塞二进制样本**：
仓库里放一个 .docx / .pdf 之后，没人说得清它里面到底有什么、是谁在什么时候生成的；
想加一个「标题 3 级」的用例就得再塞一个文件。现场生成的话，**每个用例的内容就写在用例里**，
读代码的人一眼看得到被测的输入长什么样。

生成出来的都是结构合法的真实文件（docx 是 zip + OOXML，pdf 的 xref 偏移按真实字节算），
所以它们会真的走一遍 python-docx / pdfplumber 的解析，不是假的替身。

本模块不以 `test_` 开头，pytest 不会收集它——它只是给测试用的工具箱。
"""

import io

import docx
from docx.shared import Pt

# docx 块类型。取值直接对应要验证的样式，名字就是 Word 里的样式名。
TITLE = "title"
HEADING1 = "h1"
HEADING2 = "h2"
BODY = "body"
TABLE = "table"


def build_docx(blocks: list[tuple[str, object]]) -> bytes:
    """按顺序生成一份 docx。

    blocks 里每一项是 `(类型, 内容)`：
    - `(TITLE, "文档标题")` / `(HEADING1, "小节")` / `(HEADING2, "子节")`
    - `(BODY, "正文段落")`
    - `(TABLE, [["表头", "说明"], ["net_amount", "实付金额"]])`
    """
    document = docx.Document()

    for kind, payload in blocks:
        if kind == TITLE:
            # level=0 就是 Word 的 Title 样式
            document.add_heading(str(payload), level=0)
        elif kind == HEADING1:
            document.add_heading(str(payload), level=1)
        elif kind == HEADING2:
            document.add_heading(str(payload), level=2)
        elif kind == BODY:
            document.add_paragraph(str(payload))
        elif kind == TABLE:
            rows = list(payload)  # type: ignore[arg-type]
            table = document.add_table(rows=len(rows), cols=len(rows[0]))
            for row_index, row in enumerate(rows):
                for column_index, cell_text in enumerate(row):
                    table.cell(row_index, column_index).text = str(cell_text)
        else:
            raise ValueError(f"未知的块类型：{kind!r}")

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def build_pdf(lines: list[tuple[float, float, str]]) -> bytes:
    """按顺序生成一份单页 PDF。

    lines 里每一项是 `(字号, 与下一行的基线间距, 文本)`。

    **间距必须由调用方给**，不能像正文那样自动算：换段的判定逻辑（`_normal_gap`
    与 `_is_paragraph_break`）正是被测对象，如果生成器自己按某个规则排版，
    就等于让被测代码和测试夹具共用同一套假设，测出来什么都证明不了。
    """
    script: list[str] = []
    y = 750.0

    for size, gap_after, text in lines:
        escaped = text.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")
        # 用 Tm（绝对文本矩阵）而不是 Td（相对位移）：Td 会累积上一次的位置，
        # 生成出来的坐标全乱，pdfplumber 读到的 top 也就没有意义。
        script.append(f"/F1 {size} Tf")
        script.append(f"1 0 0 1 72 {y:.2f} Tm")
        script.append(f"({escaped}) Tj")
        y -= gap_after

    stream = "\n".join(["BT", *script, "ET"]).encode("latin-1")

    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>"
        ),
        b"<< /Length "
        + str(len(stream)).encode()
        + b" >>\nstream\n"
        + stream
        + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]

    out = bytearray(b"%PDF-1.4\n")
    offsets: list[int] = []
    for index, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{index} 0 obj\n".encode() + body + b"\nendobj\n"

    xref_at = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode() + b"0000000000 65535 f \n"
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref_at}\n%%EOF\n"
    ).encode()
    return bytes(out)


def build_docx_with_font_sized_paragraphs() -> bytes:
    """一个用来验证「字号不影响 docx 判定」的辅助文档（样式才是唯一依据）。"""
    document = docx.Document()
    paragraph = document.add_paragraph("看起来像标题，但样式是正文")
    for run in paragraph.runs:
        run.font.size = Pt(28)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
